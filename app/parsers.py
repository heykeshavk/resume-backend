from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ParseError(ValueError):
    """Raised when a resume file cannot be read."""


def extract_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ParseError(
            f"Unsupported file type '{suffix}'. Upload a PDF, DOCX, or TXT resume."
        )
    if suffix == ".pdf":
        text = _from_pdf(data)
    elif suffix == ".docx":
        text = _from_docx(data)
    else:
        text = data.decode("utf-8", errors="ignore")
    text = _normalize_whitespace(text)
    if len(text) < 40:
        raise ParseError(
            "Could not extract enough text from the resume. Try a text-based PDF or DOCX."
        )
    return text


def _from_pdf(data: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise ParseError("Could not open the PDF. It may be damaged or password-protected.") from exc
    pages: list[str] = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _from_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ParseError("Could not open the Word document.") from exc
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _normalize_whitespace(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
